#!/usr/bin/env python3
"""Export the DAPHNE intake workbook to stable CSV files and verify drift."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

DAPHNE_DIR = Path(__file__).resolve().parents[1]
WORKBOOK = DAPHNE_DIR / "DAPHNE_DCS_Full_Variable_Intake.xlsx"
POLICY = DAPHNE_DIR / "DAPHNE_OPCUA_Control_Policy.csv"
EXPORT_DIR = DAPHNE_DIR / "exports"
CHECKSUMS = DAPHNE_DIR / "SHA256SUMS"
POLICY_SHEET = "OPC-UA Control Policy"
SUPPLEMENTAL_EXPORTS = ("protobuf_field_trace.csv",)
NODE_ID_PREFIX = "nsu=urn:dune:pds:daphne;s="
BOARD_PREFIX = "DAPHNE.Boards.{BoardId}."
EXTERNAL_SUBSYSTEMS = {"Authority", "Endpoint Status", "OPC-UA Bridge"}
TRACE_COLUMNS = (
    "protobuf_field_number",
    "protobuf_field",
    "field_status",
    "protobuf_type",
    "repeated",
    "instance_fields",
    "opcua_node_pattern",
    "dcs_data_type",
    "engineering_unit",
    "data_source",
    "control_owner",
)


def render(value: object) -> str:
    """Return a stable, human-readable CSV representation of an Excel value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value.replace("\r\n", "\n").replace("\r", "\n").rstrip()
    return str(value)


def filename_for_sheet(title: str) -> str:
    normalized = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "_", normalized.lower()).strip("_")
    if not slug:
        raise ValueError(f"worksheet name has no usable filename: {title!r}")
    return f"{slug}.csv"


def export_bytes() -> dict[str, bytes]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    outputs: dict[str, bytes] = {}
    try:
        for worksheet in workbook.worksheets:
            buffer = io.StringIO(newline="")
            writer = csv.writer(buffer, lineterminator="\n")
            for row in worksheet.iter_rows():
                writer.writerow([render(cell.value) for cell in row])
            name = filename_for_sheet(worksheet.title)
            if name in outputs:
                raise ValueError(f"worksheet filename collision: {name}")
            outputs[name] = buffer.getvalue().encode("utf-8")
    finally:
        workbook.close()
    return outputs


def validate_protobuf_trace(exports: dict[str, bytes]) -> list[str]:
    """Check the explicit wire-field ledger against the workbook tag export."""
    trace_path = EXPORT_DIR / "protobuf_field_trace.csv"
    if not trace_path.is_file():
        return ["missing supplemental export: exports/protobuf_field_trace.csv"]

    with trace_path.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        if tuple(reader.fieldnames or ()) != TRACE_COLUMNS:
            return ["protobuf field trace has unexpected columns"]
        trace_rows = list(reader)

    tag_rows = list(
        csv.DictReader(io.StringIO(exports["tag_list.csv"].decode("utf-8-sig")))
    )
    expected_rows = [
        row
        for row in tag_rows
        if row["DCS Access"] == "Read-Only"
        and row["Subsystem"] not in EXTERNAL_SUBSYSTEMS
    ]
    expected_by_pattern = {
        row["OPC-UA Node Address"][len(NODE_ID_PREFIX) :]: row
        for row in expected_rows
        if row["OPC-UA Node Address"].startswith(NODE_ID_PREFIX + BOARD_PREFIX)
    }
    trace_by_pattern = {row["opcua_node_pattern"]: row for row in trace_rows}
    active_trace_by_pattern = {
        pattern: row
        for pattern, row in trace_by_pattern.items()
        if row["field_status"] == "Active"
    }
    problems: list[str] = []
    if len(trace_rows) != len(trace_by_pattern):
        problems.append("protobuf field trace contains duplicate NodeId patterns")
    if any(row["field_status"] not in {"Active", "Retired"} for row in trace_rows):
        problems.append("protobuf field trace contains an invalid field status")
    if set(active_trace_by_pattern) != set(expected_by_pattern):
        problems.append("active protobuf fields differ from board-owned tag rows")

    field_numbers: list[int] = []
    field_names: list[str] = []
    for pattern, trace_row in trace_by_pattern.items():
        try:
            field_numbers.append(int(trace_row["protobuf_field_number"]))
        except ValueError:
            problems.append(f"invalid protobuf field number for {pattern}")
        field_names.append(trace_row["protobuf_field"])
        if trace_row["field_status"] != "Active":
            continue
        tag_row = expected_by_pattern.get(pattern)
        if tag_row is None:
            continue
        comparisons = (
            ("dcs_data_type", "Data Type"),
            ("engineering_unit", "Eng Units"),
            ("control_owner", "Control Owner"),
        )
        for trace_column, tag_column in comparisons:
            if trace_row[trace_column] != tag_row[tag_column]:
                problems.append(f"protobuf trace {trace_column} differs for {pattern}")
    if len(field_numbers) != len(set(field_numbers)):
        problems.append("protobuf field trace contains duplicate field numbers")
    if len(field_names) != len(set(field_names)):
        problems.append("protobuf field trace contains duplicate field names")
    return problems


def checksum_bytes(exports: dict[str, bytes]) -> bytes:
    artifacts: dict[str, bytes] = {
        WORKBOOK.name: WORKBOOK.read_bytes(),
        POLICY.name: exports[filename_for_sheet(POLICY_SHEET)],
    }
    artifacts.update({f"exports/{name}": data for name, data in exports.items()})
    artifacts.update(
        {
            f"exports/{name}": (EXPORT_DIR / name).read_bytes()
            for name in SUPPLEMENTAL_EXPORTS
        }
    )
    lines = [
        f"{hashlib.sha256(data).hexdigest()}  {name}\n"
        for name, data in sorted(artifacts.items())
    ]
    return "".join(lines).encode("ascii")


def update() -> None:
    exports = export_bytes()
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    for name, data in exports.items():
        (EXPORT_DIR / name).write_bytes(data)
    POLICY.write_bytes(exports[filename_for_sheet(POLICY_SHEET)])
    problems = validate_protobuf_trace(exports)
    if problems:
        raise ValueError("; ".join(problems))
    CHECKSUMS.write_bytes(checksum_bytes(exports))
    print(
        f"exported {len(exports)} worksheets, verified explicit protobuf trace, "
        f"and refreshed {CHECKSUMS.name}"
    )


def check() -> int:
    exports = export_bytes()
    problems: list[str] = []
    expected_names = set(exports) | set(SUPPLEMENTAL_EXPORTS)
    actual_names = {path.name for path in EXPORT_DIR.glob("*.csv")}

    for name in sorted(expected_names - actual_names):
        problems.append(f"missing export: exports/{name}")
    for name in sorted(actual_names - expected_names):
        problems.append(f"unexpected export: exports/{name}")
    for name in sorted(expected_names & actual_names):
        if name in exports and (EXPORT_DIR / name).read_bytes() != exports[name]:
            problems.append(f"stale export: exports/{name}")

    problems.extend(validate_protobuf_trace(exports))

    policy_data = exports[filename_for_sheet(POLICY_SHEET)]
    if not POLICY.exists() or POLICY.read_bytes() != policy_data:
        problems.append(f"stale policy export: {POLICY.name}")

    if all((EXPORT_DIR / name).is_file() for name in SUPPLEMENTAL_EXPORTS):
        expected_checksums = checksum_bytes(exports)
        if not CHECKSUMS.exists() or CHECKSUMS.read_bytes() != expected_checksums:
            problems.append(f"stale integrity file: {CHECKSUMS.name}")

    if problems:
        for problem in problems:
            print(problem, file=sys.stderr)
        return 1

    print(f"verified {len(exports)} worksheet exports, policy, and checksums")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--check",
        action="store_true",
        help="report drift without rewriting generated files",
    )
    args = parser.parse_args()
    if not WORKBOOK.is_file():
        parser.error(f"workbook not found: {WORKBOOK}")
    return check() if args.check else (update() or 0)


if __name__ == "__main__":
    raise SystemExit(main())
